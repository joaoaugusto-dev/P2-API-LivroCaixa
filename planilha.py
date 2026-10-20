"""Exportação para Excel (.xlsx) com a identidade visual do Livro-caixa. Usa a biblioteca openpyxl."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

# Paleta do Livro-caixa (a mesma do styles.css), para a planilha ter a identidade do sistema
PETROLEO, PETROLEO_2, OURO = "0D3A40", "164A52", "F2B632"
ENTRADA, SAIDA, FUNDO, LINHA, SUAVE = "14805F", "C8472F", "EEF3F0", "D3DCD8", "566A6E"
MOEDA = '"R$" #,##0.00;-"R$" #,##0.00'
FONTE = "Calibri"


# Atalho: toda fonte da planilha usa Calibri
def _fonte(**kw):
    return Font(name=FONTE, **kw)


def gerar_xlsx(linhas, filtro):
    """linhas: o mesmo formato de db.listar(). filtro: texto que descreve o recorte."""
    # 1) Cria a pasta de trabalho e configura a aba (cor, sem linhas de grade, larguras)
    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato"
    ws.sheet_properties.tabColor = PETROLEO
    ws.sheet_view.showGridLines = False
    for col, largura in zip("ABCD", (14, 52, 18, 4)):
        ws.column_dimensions[col].width = largura

    # 2) Faixa de título verde-petróleo com um filete dourado embaixo
    for r in (1, 2):
        for c in "ABC":
            ws[f"{c}{r}"].fill = PatternFill("solid", fgColor=PETROLEO)
    ws["A1"] = "Livro-caixa"
    ws["A1"].font = _fonte(size=20, bold=True, color="FFFFFF")
    ws["A2"] = filtro
    ws["A2"].font = _fonte(size=10, color="9FBBB8")
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    for c in "ABC":
        ws[f"{c}3"].fill = PatternFill("solid", fgColor=OURO)
    ws.row_dimensions[3].height = 4

    # 3) Indicadores (entradas, saídas, saldo) como fórmulas: recalculam se a pessoa editar a planilha
    topo, fim = 10, 10 + max(len(linhas), 1) - 1
    faixa_valor = f"$C${topo}:$C${fim}"
    kpis = [
        ("Entradas", f'=SUMIF({faixa_valor},">0")', ENTRADA),
        ("Saídas", f'=SUMIF({faixa_valor},"<0")', SAIDA),
        ("Saldo", f"=SUM({faixa_valor})", PETROLEO),
    ]
    for i, (nome, formula, cor) in enumerate(kpis):
        ws[f"A{5 + i}"] = nome
        ws[f"A{5 + i}"].font = _fonte(size=11, color=SUAVE)
        ws.merge_cells(f"A{5 + i}:B{5 + i}")
        ws[f"C{5 + i}"] = formula
        ws[f"C{5 + i}"].number_format = MOEDA
        ws[f"C{5 + i}"].font = _fonte(size=13 if nome != "Saldo" else 15, bold=True, color=cor)
        ws[f"C{5 + i}"].alignment = Alignment(horizontal="right")
        ws[f"A{5 + i}"].border = ws[f"C{5 + i}"].border = Border(bottom=Side(style="thin", color=LINHA))
        ws[f"B{5 + i}"].border = Border(bottom=Side(style="thin", color=LINHA))
    ws["A8"] = f"{len(linhas)} movimentações"
    ws["A8"].font = _fonte(size=10, italic=True, color=SUAVE)

    # 4) Tabela: cabeçalho colorido + uma linha por movimentação
    cab = topo - 1
    for c, titulo in zip("ABC", ("Data", "Descrição", "Valor")):
        cel = ws[f"{c}{cab}"]
        cel.value = titulo
        cel.font = _fonte(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=PETROLEO_2)
        cel.alignment = Alignment(horizontal="right" if c == "C" else "left", vertical="center", indent=1)
    ws.row_dimensions[cab].height = 24

    for i, (_, data, valor, desc, _) in enumerate(linhas):
        r = topo + i
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = data, desc, float(valor)
        ws[f"A{r}"].number_format = "DD/MM/YYYY"
        ws[f"A{r}"].font = _fonte(color=SUAVE)
        ws[f"B{r}"].font = _fonte()
        ws[f"C{r}"].number_format = MOEDA
        ws[f"C{r}"].font = _fonte(bold=True, color=ENTRADA if valor > 0 else SAIDA)
        for c in "ABC":
            cel = ws[f"{c}{r}"]
            cel.border = Border(bottom=Side(style="hair", color=LINHA))
            cel.alignment = Alignment(horizontal="right" if c == "C" else "left", indent=1, vertical="center")
            if i % 2:
                cel.fill = PatternFill("solid", fgColor=FUNDO)
        ws.row_dimensions[r].height = 20

    # 5) Acabamento: filtro, cabeçalho congelado e impressão ajustada à largura da página
    if linhas:
        ws.auto_filter.ref = f"A{cab}:C{fim}"
    ws.freeze_panes = f"A{topo}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{cab}:{cab}"

    # 6) Salva em memória e devolve os bytes (a API manda direto como download)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
